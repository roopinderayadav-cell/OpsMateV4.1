"""OpsMate checks.  Run:  pip install pytest  then  pytest -q test_opsmate.py"""
import io
import zipfile

import pytest

import config
import knowledge as kb


@pytest.fixture(autouse=True)
def fresh(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "knowledge.db")
    monkeypatch.setattr(kb.llm, "ai_on", lambda: False)
    monkeypatch.setattr(kb.llm, "embeddings_on", lambda: False)
    kb.init_db()
    kb.load_samples_if_empty()


def test_lob_tabs_and_sub_lobs():
    lobs = kb.get_lobs()
    assert list(lobs) == ["NORAM", "EUROPE", "APAC", "QBAY"]
    assert "UETTR" in lobs["EUROPE"] and "HK Refund" in lobs["APAC"] and lobs["QBAY"] == ["User Guide"]
    kb.save_lobs({"QBAY": ["User Guide", " user guide ", "", "Release Notes"]})
    assert kb.get_lobs()["QBAY"] == ["User Guide", "Release Notes"]
    assert "UETTR" in kb.get_lobs()["EUROPE"]                 # other LOBs untouched
    with pytest.raises(ValueError):
        kb.save_lobs({"LATAM": ["x"]})


def test_sub_lob_documents_rank_first():
    docs = {d["title"]: d for d in kb.list_documents()}
    assert docs["OBT Cancellation SOP"]["sub_lob"] == "OBT Cancellation"
    a = kb.ask("NORAM", "how to process refund", "t", sub_lob="OBT Cancellation")
    assert a["result"] == "answered" and a["sources"][0]["sub_lob"] == "OBT Cancellation"
    b = kb.ask("NORAM", "how to process refund", "t", sub_lob="Refund")
    assert b["sources"][0]["sub_lob"] == "Refund"
    c = kb.ask("NORAM", "what is TTP", "t", sub_lob="OBT Cancellation")     # other sub-LOB still searched
    assert c["result"] == "answered" and "TTP" in c["answer"]


def test_other_lob_fallback_is_labelled_and_never_high():
    a = kb.ask("EUROPE", "traveller cannot log in to the obt", "t")   # only a NORAM document answers this
    assert a["result"] == "answered"
    assert a["meta"]["other_lob"] == "NORAM" and a["confidence"] != "high"
    assert any("Not found in EUROPE" in n for n in a["meta"]["notes"])


def test_confidence_percent_matches_level_band():
    for q in ["how to procss refund", "what is TTP", "list of amadeus commands", "can i void tiket next day"]:
        a = kb.ask("NORAM", q, "t")
        lo, hi = config.CONF_BANDS[a["confidence"]]
        assert lo <= a["confidence_pct"] <= hi
    assert kb.ask("NORAM", "what is the canteen menu today", "t")["confidence_pct"] is None


def test_answers_are_steps_for_how_to_questions():
    a = kb.ask("NORAM", "refund of self booked ticket", "t", sub_lob="OBT Cancellation")
    assert a["answer"].lstrip().startswith("1.") and "2." in a["answer"]


def test_backup_restore_roundtrip_and_rejects_bad_zip():
    before = len(kb.list_documents())
    data = kb.make_backup()
    kb.delete_document(kb.list_documents()[0]["id"])
    assert len(kb.list_documents()) == before - 1
    assert kb.restore_backup(data) == before
    assert len(kb.list_documents()) == before
    assert all(kb.read_file(d) for d in kb.list_documents())
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("knowledge.db", b"x")
        z.writestr("uploads/../../evil.py", b"x")
    with pytest.raises(ValueError):
        kb.restore_backup(bad.getvalue())
    with pytest.raises(ValueError):
        kb.restore_backup(b"not a zip")


def _docx_with_picture():
    """A small Word SOP: one Amadeus and one Sabre section, a sample screen picture in the Amadeus one."""
    import docx
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (640, 260), "#0f2238")
    d = ImageDraw.Draw(img)
    for k in range(9):
        d.text((20, 20 + 24 * k), f"{k + 1} SAMPLE LINE {k * 37:05d} RP/DELXX2100", fill="#7CFC9A")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = docx.Document()
    doc.add_heading("Amadeus - How to void a ticket step by step", 2)
    for t in ["1. Retrieve the PNR and display the ticket: TWD.", "2. Check the ticket is OPEN and issued today.",
              "3. Void the ticket: TWX/L5. Check the coupon status is VOID."]:
        doc.add_paragraph(t)
    buf.seek(0)
    doc.add_picture(buf)
    doc.add_heading("Sabre - How to void a ticket step by step", 2)
    for t in ["1. Retrieve the PNR: *ABCDEF and display the ticket field: *T.",
              "2. Check the ticket is OPEN and issued today: WETR*1.", "3. Void the ticket: WV1. Check the status is VOID."]:
        doc.add_paragraph(t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_pictures_are_stored_and_shown_with_the_answer():
    doc_id, _, info = kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    assert info["images_stored"] == 1
    r = kb.ask("APAC", "how to void a ticket in amadeus", "tester")
    assert r["result"] == "answered" and len(r["images"]) == 1
    assert "Amadeus - How to void a ticket" in r["images"][0]["caption"]
    s = kb.ask("APAC", "how to void a ticket in sabre", "tester")
    assert s["result"] == "answered" and s["images"] == []          # Sabre section has no picture
    kb.delete_document(doc_id)
    assert not list(config.UPLOAD_DIR.glob(f"docimg_{doc_id}_*"))     # pictures removed with the document


def test_backfill_keeps_pictures_of_older_documents():
    doc_id, _, _ = kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    kb._delete_images(doc_id)
    with kb._conn() as con:
        con.execute("UPDATE documents SET images_done=0 WHERE id=?", (doc_id,))
    assert kb.backfill_images() == 1 and kb.backfill_images() == 0  # second start: nothing to do


def test_asks_which_gds_when_not_said_and_answers_after_choice():
    kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    r = kb.ask("APAC", "how to void a ticket", "tester")
    assert r["result"] == "clarify" and r["id"] is None
    assert {o["label"] for o in r["meta"]["clarify"]["options"]} == {"Amadeus", "Sabre"}
    a = kb.ask("APAC", "how to void a ticket in Sabre", "tester", clarify=False)
    assert a["result"] == "answered" and "WV1" in a["answer"]
    assert kb.ask("APAC", "how to void a ticket", "tester", clarify=False)["result"] == "answered"


def test_gds_commands_are_detected_without_false_positives():
    import brain
    t = "Retrieve RTABC123, price FXP, issue TTP/RT, Sabre *T and WETR*1. APAC ADM NON-REF AI101 DEL-BOM HK1"
    assert brain.gds_commands(t) == ["RTABC123", "FXP", "TTP/RT", "*T", "WETR*1"]
    assert "`TTP/RT`" in brain.highlight_commands(t) and "`APAC`" not in brain.highlight_commands(t)


def test_document_answer_when_ai_says_not_found(monkeypatch):
    """Busy free tier: the AI replies 'not found' although the SOP clearly answers - show the SOP answer."""
    monkeypatch.setattr(kb.llm, "ai_on", lambda: True)
    monkeypatch.setattr(kb.brain, "ai_answer", lambda *a, **k: {
        "status": "not_found", "confidence": "low", "answer": "", "important": "", "conflict": "", "used": []})
    monkeypatch.setattr(kb.brain, "ai_understand", lambda *a, **k: {"corrected": a[0], "queries": [], "format": "short"})
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]
    assert any("taken directly from the documents" in n for n in r["meta"]["notes"])
    assert kb.ask("NORAM", "what is the canteen menu today", "tester")["result"] == "not_found"   # still refuses


def test_quota_error_pauses_ai_and_answers_from_documents(monkeypatch):
    import time
    calls = {"n": 0}

    def quota(*a, **k):
        calls["n"] += 1
        raise RuntimeError("429 RESOURCE_EXHAUSTED: quota exceeded")
    monkeypatch.setattr(kb.llm, "provider", lambda: "gemini")
    monkeypatch.setattr(kb.llm, "ai_on", lambda: True)
    monkeypatch.setattr(kb.llm, "_generate", lambda *a, **k: kb.llm._retry(quota))
    monkeypatch.setattr(kb.llm.time, "sleep", lambda s: None)
    monkeypatch.setattr(kb.llm, "_paused_until", {"generate": 0.0, "embed": 0.0})
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]
    assert kb.llm.paused("generate") and calls["n"] >= 1
    before = calls["n"]
    t0 = time.time()
    r2 = kb.ask("NORAM", "what is TST", "tester")
    assert r2["result"] == "answered" and calls["n"] == before        # AI skipped while paused
    assert time.time() - t0 < 5


def test_generic_words_do_not_block_an_answer():
    import brain
    sec = "Sabre - Availability and schedules\nAvailability for a date and city pair: 1 + date + from + to"
    assert brain.coverage("how to check the availability in sabre GDS", sec) == 1.0
    assert brain.coverage("what is the command to print", "Sabre - Mandatory PNR fields (PRINT)") < 0.7
    assert kb.definition_term("what is AGM process") == "agm" and kb.definition_term("what is the BRIM system?") == "brim"


def test_follow_up_questions_use_the_previous_question():
    import brain
    assert brain.merge_followup("How to reissue a ticket in Amadeus?", "how about Sabre") == "How to reissue a ticket in Sabre?"
    assert brain.merge_followup("What is AGM?", "and BRIM?", ["AGM", "BRIM"]) == "What is BRIM?"
    assert brain.merge_followup("How to void a ticket?", "what is ADM") == "what is ADM"       # a new question


def test_try_asking_follows_the_selected_lob():
    noram = kb.suggested_questions("NORAM", None)
    assert noram and len(noram) <= 6
    obt = kb.suggested_questions("NORAM", "OBT Cancellation")
    assert obt and obt != noram


def test_everyday_words_are_not_misread():
    import brain
    vocab = {"gives", "shared", "hotel", "invoice", "refund", "ticket"}
    assert brain.fix_spelling("give me hotel invoice", vocab)[0] == "give me hotel invoice"   # not 'gives'
    assert brain.fix_spelling("refnd tiket", vocab)[0] == "refund ticket"                     # real typos still fixed
    assert "looks" not in brain.search_text("how does the hotel invoice looks in Europe")
    sec = "Hotel invoice (Norway) - fields and what to check\nThe supplier who issued the invoice"
    assert brain.coverage("do you have any hotel invoice image which i refer", sec) == 1.0
